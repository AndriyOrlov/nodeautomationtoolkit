"""Таблиця з даними про осіб: графи впізнаються за заголовком, а не за номером.

План звільнення (додаток 21) і списки до присвоєння звань у кожній частині
роблять по-своєму, тому колонки читаються за назвами з довідника граф.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from nodeautomationtoolkit.order_generator import pipeline  # noqa: E402
from nodeautomationtoolkit.order_generator.compose import DISMISSAL, OrderParams  # noqa: E402
from nodeautomationtoolkit.order_generator.table_plan import read_table  # noqa: E402

DISMISSAL_HEADER = [
    "№ з/п",
    "Військове звання",
    "Прізвище, ім'я та по батькові",
    "РНОКПП",
    "Займана посада",
    "Дата народження",
    "Вислуга календарна",
    "Вислуга пільгова",
    "Підстава звільнення",
    "Куди звільняється",
    "Військовий облік",
    "Право на форму",
]

DISMISSAL_ROWS = [
    [
        "1",
        "полковник",
        "БОНДАР Руслан Володимирович",
        "2652323639",
        "заступник командира батальйону",
        "11.08.1976",
        "29 років 3 місяці",
        "29 років 11 місяців",
        "1.а",
        "у запас",
        "Слобідського ОРТЦК та СП м. Харкова",
        "так",
    ],
    [
        "2",
        "підполковник",
        "АНТОНЮК Віталій Віталійович",
        "2496955999",
        "начальник служби",
        "12.05.1968",
        "35 років 4 місяці",
        "41 рік 6 місяців",
        "1.б",
        "у відставку",
        "Голосіївського районного у місті Києві ТЦК та СП",
        "так",
    ],
]


@pytest.fixture
def dismissal_plan(tmp_path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ЗАТВЕРДЖУЮ"])
    sheet.append(["ПЛАН ЗВІЛЬНЕННЯ осіб офіцерського складу"])
    sheet.append(DISMISSAL_HEADER)
    for row in DISMISSAL_ROWS:
        sheet.append(row)
    path = tmp_path / "План звільнення.xlsx"
    workbook.save(path)
    return path


def test_columns_are_found_by_their_titles(dismissal_plan):
    plan = read_table(dismissal_plan)
    assert plan.header_row == 3
    assert len(plan.records) == 2
    assert not plan.problems

    first = plan.records[0]
    assert first.full_name == "БОНДАР Руслан Володимирович"
    assert str(first.rank) == "полковник"
    assert str(first.ipn) == "2652323639"
    assert str(first.current.text) == "заступник командира батальйону"
    assert str(first.birth) == "11.08.1976"
    assert str(first.service_calendar) == "29 років 3 місяці"
    assert str(first.dismissal) == "1.а"
    assert str(first.destination) == "у запас"
    assert str(first.registration).startswith("Слобідського")
    assert str(first.uniform) == "так"


def test_unknown_columns_are_reported(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Військове звання", "Прізвище, ім'я та по батькові", "РНОКПП", "Хитра графа"])
    sheet.append(["майор", "ІВАНЕНКО Олексій Вікторович", "3508402997", "щось своє"])
    path = tmp_path / "список.xlsx"
    workbook.save(path)

    plan = read_table(path)
    assert [record.full_name for record in plan.records] == ["ІВАНЕНКО Олексій Вікторович"]
    assert "Хитра графа" in plan.unknown_columns


def test_a_table_without_known_columns_says_so(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["раз", "два", "три"])
    sheet.append(["1", "2", "3"])
    path = tmp_path / "чужа.xlsx"
    workbook.save(path)

    plan = read_table(path)
    assert not plan.records
    assert any("довідник граф" in problem for problem in plan.problems)


def test_the_column_order_does_not_matter(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["РНОКПП", "Займана посада", "Прізвище, ім'я та по батькові", "Звання"])
    sheet.append(["2652323639", "командир роти", "БОНДАР Руслан Володимирович", "капітан"])
    path = tmp_path / "інший порядок.xlsx"
    workbook.save(path)

    record = read_table(path).records[0]
    assert str(record.rank) == "капітан"
    assert str(record.current.text) == "командир роти"
    assert str(record.ipn) == "2652323639"


def test_dismissal_order_straight_from_the_table(dismissal_plan, tmp_path):
    records = pipeline.records_from_table(dismissal_plan, action=DISMISSAL)
    result = pipeline.run(
        records,
        OrderParams(
            action=DISMISSAL,
            law_points="пункту другого частини п'ятої статті 26",
            number="530",
            date="18.09.2026",
        ),
        tmp_path / "готове",
    )
    assert result.document is not None and result.document.exists()

    from docx import Document

    text = "\n".join(p.text for p in Document(str(result.document)).paragraphs).replace(" ", " ")
    assert "У ЗАПАС ЗА ПІДПУНКТОМ «а»" in text
    assert "У ВІДСТАВКУ ЗА ПІДПУНКТОМ «б»" in text
    assert "Полковника БОНДАРА Руслана Володимировича" in text
    assert "Народився 11 серпня 1976 року" in text
