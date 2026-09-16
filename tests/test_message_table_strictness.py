"""Шифр у закритому змісті — СТРОГО з таблиці, нові частини видно (AGENT.md 9.5.7).

Закріплює дефекти, знайдені перевіркою генерації повідомлень (усі дані вигадані):

1. Порожній стовпець B: `read_recipient_mapping` кладе в `cipher` відкриту
   назву, і в закритий текст ішло «військової частини 77 окрема тестова бригада».
2. Корпус зі стовпця D, якого немає в таблиці: у текст ішов сам стовпець D —
   «військової частини А0077 військової частини 99 АК».
3. Описки виправлялися в САМОМУ тексті: абзац без жодної частини змінювався
   (`зв’язку` → `зв'язку`, `в. ч.` → `в/ч`).
4. Нові частини («169 батальйону резерву», «12 навчального центру») лишались
   відкритими без жовтої позначки.
"""

import openpyxl
import pytest

import generate_extracts as generator
from nodeautomationtoolkit.builtin_nodes.message_order import cipher_unit_names
from nodeautomationtoolkit.builtin_nodes.recipient_mapping import (
    _format_full_closed_unit_text,
    map_military_units,
    read_recipient_mapping,
)

BRIGADE = "77 окрема тестова бригада"
HEADER = ["Відкрите найменування", "Шифр", "Скорочення", "Корпус", "Кому", "Куди"]


def _entry(open_name, cipher, abbreviation="", corps="", recipient_to=""):
    return {
        "open_name": open_name,
        "cipher": cipher,
        "abbreviation": abbreviation,
        "corps": corps,
        "recipient_to": recipient_to,
    }


def _xlsx(tmp_path, rows, name="словник.xlsx"):
    path = tmp_path / name
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(HEADER)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


class _Var:
    def __init__(self, value):
        self.value = value

    def get(self):
        return self.value


# ── 1–2. Шифр лише з таблиці ────────────────────────────────────────────────


def test_empty_column_b_leaves_name_open_and_reports(tmp_path):
    table = tmp_path / "словник.csv"
    table.write_text(";".join(HEADER) + f"\n{BRIGADE};;77 отбр;;;\n", encoding="utf-8")
    mapping = read_recipient_mapping(str(table))["mapping"]
    text = "командира взводу 77 окремої тестової бригади"

    problems = []
    result, count, rows = cipher_unit_names(text, mapping, problems=problems)

    assert result == text
    assert count == 0
    assert rows == []
    assert problems == [("no_cipher", BRIGADE)]


def test_unknown_corps_in_column_d_adds_no_invented_link():
    mapping = {BRIGADE: _entry(BRIGADE, "А0077", "77 отбр", "99 АК")}

    problems = []
    result, _, _ = cipher_unit_names("командира взводу 77 окремої тестової бригади", mapping, problems=problems)

    assert result == "командира взводу військової частини А0077"
    assert problems == [("corps_missing", BRIGADE, "99 АК")]


def test_corps_row_without_cipher_adds_no_link():
    corps = "51 армійський корпус"
    mapping = {
        BRIGADE: _entry(BRIGADE, "А0077", "77 отбр", "51 АК"),
        corps: _entry(corps, corps, "51 АК"),  # B порожній — у cipher відкрита назва
    }

    result, _, _ = cipher_unit_names("командира взводу 77 окремої тестової бригади", mapping)

    assert result == "командира взводу військової частини А0077"
    assert _format_full_closed_unit_text(mapping[BRIGADE], mapping) == "військової частини А0077"


def test_problems_are_reported_once_per_row():
    mapping = {BRIGADE: _entry(BRIGADE, "А0077", "77 отбр", "99 АК")}
    problems = []
    cipher_unit_names(
        "з 77 окремої тестової бригади до 77 окремої тестової бригади", mapping, problems=problems
    )
    assert problems == [("corps_missing", BRIGADE, "99 АК")]


# ── 3. Описки — лише для пошуку ─────────────────────────────────────────────


def test_paragraph_without_units_is_not_rewritten():
    source = "начальника вузла зв’язку в. ч. тестового гарнізону"
    result, count, _ = cipher_unit_names(source, {BRIGADE: _entry(BRIGADE, "А0077")})
    assert result == source
    assert count == 0


def test_name_with_typo_is_found_but_rest_of_text_stays_original():
    mapping = {BRIGADE: _entry(BRIGADE, "А0077")}
    result, _, _ = cipher_unit_names("начальник зв’язку 77 окремої тестової бригди", mapping)
    assert result == "начальник зв’язку військової частини А0077"


# ── 4. Жовта позначка для нових частин ──────────────────────────────────────


@pytest.mark.parametrize(
    "text, expected",
    [
        ("офіцера 169 батальйону резерву", "169 батальйону"),
        ("офіцера 12 навчального центру", "12 навчального центру"),
        ("офіцера 300 військового госпіталю", "300 військового госпіталю"),
        ("офіцера 40 бази зберігання", "40 бази"),
        ("офіцера 88 окремого вузла зв'язку", "88 окремого вузла"),
    ],
)
def test_numbered_units_of_any_kind_are_highlighted(text, expected):
    spans = generator.find_unmatched_open_unit_spans(text)
    assert [text[start:end] for start, end in spans] == [expected]


@pytest.mark.parametrize(
    "text",
    [
        "офіцера 2 відділу Калуського районного територіального центру комплектування",
        "командира 1 механізованого батальйону цієї самої військової частини",
        "командира 2 батальйону військової частини А1111",
        "з 01.06.2026 року",
        "відповідно до пункту 2 частини четвертої статті 26 Закону України",
    ],
)
def test_subunits_dates_and_references_are_not_highlighted(text):
    assert generator.find_unmatched_open_unit_spans(text) == []


def test_new_units_are_collected_once_with_full_name():
    mapping = {
        BRIGADE: _entry(BRIGADE, "А0077"),
        "55 окремий батальйон тестування": _entry(
            "55 окремий батальйон тестування", "55 окремий батальйон тестування"
        ),
    }
    text = (
        "1. Лейтенанта ТЕСТЕНКА Теста Тестовича, командира взводу 77 окремої тестової бригади, "
        "офіцера 169 батальйону резерву.\n"
        "2. Капітана ВИГАДАНОГО Вигада Вигадовича, офіцера 55 окремого батальйону тестування.\n"
        "3. Майора ПРИКЛАДОВА Приклада Прикладовича, офіцера 169 батальйону резерву."
    )

    # Частина з порожнім B у таблиці вже є — нова лише одна, і без повторів.
    assert generator.collect_new_unit_names(text, mapping) == ["169 батальйону резерву"]


def test_recipient_without_cipher_does_not_leak_open_name():
    mapping = {BRIGADE: _entry(BRIGADE, BRIGADE, "77 отбр", "", "Командиру 77 отбр")}
    routes = map_military_units(text="§ 1\n1. Направити до 77 окремої тестової бригади.", mapping=mapping)

    recipients = generator.build_message_recipient_list(mapping, routes)

    assert "окрема тестова бригада" not in " ".join(recipients)


# ── Заготовки нових частин у таблиці ────────────────────────────────────────


def test_stubs_are_appended_to_table_without_formulas(tmp_path):
    path = _xlsx(tmp_path, [[BRIGADE, "А0077", "77 отбр", "", "", ""]])

    result = generator.append_unit_stubs_to_table(
        str(path), ["169 батальйону резерву", "169  батальйону резерву ", BRIGADE]
    )

    assert result["added"] == ["169 батальйону резерву"]
    assert result["separate"] is False
    assert openpyxl.load_workbook(result["backup"]).active.max_row == 2
    sheet = openpyxl.load_workbook(path).active
    assert sheet.cell(row=3, column=1).value == "169 батальйону резерву"
    assert sheet.cell(row=3, column=2).value is None
    assert sheet.cell(row=3, column=1).fill.start_color.rgb.endswith("FFFF00")

    # Заготовка без шифру маршрутизації не зачіпає, а шифр частини цілий.
    mapping = read_recipient_mapping(str(path))["mapping"]
    assert "169 батальйону резерву" not in mapping
    assert mapping[BRIGADE]["cipher"] == "А0077"

    # Повторний запуск не дублює заготовку.
    assert generator.append_unit_stubs_to_table(str(path), ["169 батальйону резерву"])["added"] == []


def test_formula_table_is_written_by_excel_when_available(tmp_path, monkeypatch):
    """Таблицю з формулами дописує сам Excel: openpyxl стер би обчислені значення."""
    path = _xlsx(tmp_path, [[BRIGADE, '="А"&"0077"', "77 отбр", "", "", ""]])
    calls = []
    monkeypatch.setattr(
        generator, "_append_rows_with_excel", lambda source, names, column: calls.append((str(source), list(names), column)) or True
    )

    result = generator.append_unit_stubs_to_table(str(path), ["169 батальйону резерву"])

    assert result["added"] == ["169 батальйону резерву"]
    assert result["separate"] is False
    assert calls == [(str(path), ["169 батальйону резерву"], 1)]
    assert result["backup"] and openpyxl.load_workbook(result["backup"]).active.max_row == 2


def test_table_with_formulas_is_not_touched_without_excel(tmp_path, monkeypatch):
    """Без Excel (як на CI) таблиця лишається недоторканою, заготовки — окремо."""
    path = _xlsx(tmp_path, [[BRIGADE, '="А"&"0077"', "77 отбр", "", "", ""]])
    monkeypatch.setattr(generator, "_append_rows_with_excel", lambda *args, **kwargs: False)

    result = generator.append_unit_stubs_to_table(str(path), ["169 батальйону резерву"])

    assert result["separate"] is True
    assert result["path"].endswith("— нові частини.xlsx")
    assert openpyxl.load_workbook(path).active.max_row == 2
    stub_sheet = openpyxl.load_workbook(result["path"]).active
    assert stub_sheet.cell(row=2, column=1).value == "169 батальйону резерву"


def test_stubs_are_appended_to_csv_table(tmp_path):
    path = tmp_path / "словник.csv"
    path.write_text(";".join(HEADER) + f"\n{BRIGADE};А0077;77 отбр;;;", encoding="utf-8")

    result = generator.append_unit_stubs_to_table(str(path), ["169 батальйону резерву"])

    assert result["added"] == ["169 батальйону резерву"]
    assert path.read_text(encoding="utf-8").splitlines()[-1] == "169 батальйону резерву;;;;;"
    assert "169 батальйону резерву" not in read_recipient_mapping(str(path))["mapping"]


def test_report_table_gaps_adds_stubs_only_when_enabled(tmp_path):
    path = _xlsx(tmp_path, [[BRIGADE, "А0077", "77 отбр", "", "", ""]])
    app = generator.App.__new__(generator.App)
    logs = []
    app.log = logs.append
    app.excel_path = _Var(str(path))
    mapping = read_recipient_mapping(str(path))["mapping"]
    order = "НАКАЗ\n§ 1\n1. Лейтенанта ТЕСТЕНКА Теста Тестовича, офіцера 169 батальйону резерву.\n"

    assert app._report_table_gaps(order, mapping) == {"problems": 0, "new_units": 1}
    assert any("169 батальйону резерву" in line for line in logs)
    assert openpyxl.load_workbook(path).active.max_row == 2

    app.ADD_NEW_UNITS_TO_TABLE = True
    app._report_table_gaps(order, mapping)
    assert openpyxl.load_workbook(path).active.cell(row=3, column=1).value == "169 батальйону резерву"


# ── «КУДИ» ВЕЛИКИМИ: частини з таблиці не стають «новими» ───────────────────


_UPPER_MAPPING = {
    "14 окремий полк зв'язку": _entry("14 окремий полк зв'язку", "А1414", "14 опз"),
    "22 окремий полк забезпечення": _entry(
        "22 окремий полк забезпечення", "22 окремий полк забезпечення", "22 опз"
    ),
    "55 окремий полк радіотехнічного забезпечення": _entry(
        "55 окремий полк радіотехнічного забезпечення", "А5555", "55 опрз"
    ),
}


def test_uppercase_name_with_typographic_apostrophe_is_ciphered():
    """Виправлення описок робило «ЗВ’ЯЗКУ» малим, і збіг відкидався як мішаний."""
    text = "командира взводу – КОМАНДИРОМ РОТИ 14 ОКРЕМОГО ПОЛКУ ЗВ’ЯЗКУ."
    result, _, _ = cipher_unit_names(text, _UPPER_MAPPING)
    assert result == "командира взводу – КОМАНДИРОМ РОТИ ВІЙСЬКОВОЇ ЧАСТИНИ А1414."
    assert generator.collect_new_unit_names(text, _UPPER_MAPPING) == []


def test_uppercase_row_without_cipher_is_not_a_new_unit():
    text = "– КОМАНДИРОМ РОТИ 22 ОКРЕМОГО ПОЛКУ ЗАБЕЗПЕЧЕННЯ."
    assert generator.collect_new_unit_names(text, _UPPER_MAPPING) == []


def test_partial_mention_of_table_unit_is_reported_as_similar_not_new():
    similar = []
    names = generator.collect_new_unit_names(
        "– НАЧАЛЬНИКОМ ШТАБУ 55 ОКРЕМОГО ПОЛКУ.", _UPPER_MAPPING, similar=similar
    )
    assert names == []
    assert similar == [("55 ОКРЕМОГО ПОЛКУ", "55 окремий полк радіотехнічного забезпечення")]


def test_uppercase_new_unit_keeps_full_name():
    text = "– КОМАНДИРОМ ВЗВОДУ 169 БАТАЛЬЙОНУ РЕЗЕРВУ."
    assert generator.collect_new_unit_names(text, _UPPER_MAPPING) == ["169 БАТАЛЬЙОНУ РЕЗЕРВУ"]


# ── Пом'якшувальний фільтр: апострофи, дефіси, регістр ──────────────────────


@pytest.mark.parametrize(
    "table_name, text",
    [
        ("12 окремий Камʼянець-Подільський батальйон", "офіцера 12 окремого Кам’янець-Подільського батальйону"),
        ("12 окремий Кам’янець-Подільський батальйон", "– ОФІЦЕРОМ 12 ОКРЕМОГО КАМ'ЯНЕЦЬ–ПОДІЛЬСЬКОГО БАТАЛЬЙОНУ"),
        ("42 окрема гірсько-штурмова бригада", "командира 42 окремої гірсько - штурмової бригади"),
        ("42 окрема гірсько-штурмова бригада", "– КОМАНДИРОМ 42 ОКРЕМОЇ ГІРСЬКО - ШТУРМОВОЇ БРИГАДИ"),
        ("42 ОКРЕМА ГІРСЬКО-ШТУРМОВА БРИГАДА", "командира 42 окремої гірсько‑штурмової бригади"),
    ],
)
def test_softened_search_ignores_apostrophe_dash_and_case_variants(table_name, text):
    mapping = {table_name: _entry(table_name, "А4242")}
    result, count, _ = cipher_unit_names(text, mapping)
    assert count == 1
    assert "А4242" in result
    assert generator.collect_new_unit_names(text, mapping) == []


def test_softening_keeps_original_text_outside_names():
    mapping = {"42 окрема гірсько-штурмова бригада": _entry("42 окрема гірсько-штурмова бригада", "А4242")}
    text = "командира зв’язку десантно - штурмового взводу 42 окремої гірсько - штурмової бригади"
    result, _, _ = cipher_unit_names(text, mapping)
    assert result == "командира зв’язку десантно - штурмового взводу військової частини А4242"


def test_spaced_dash_between_source_and_destination_is_not_joined():
    """Межа «звідки – КУДИ» (мале → ВЕЛИКЕ) не склеюється — інакше зникав би текст."""
    mapping = {"42 окрема гірсько-штурмова бригада": _entry("42 окрема гірсько-штурмова бригада", "А4242")}
    text = "командира взводу 42 окремої гірсько - ШТУРМОВОЇ БРИГАДИ"
    result, count, _ = cipher_unit_names(text, mapping)
    assert result == text
    assert count == 0


# ── Почесні найменування (форми з додатку 53, назви вигадані) ───────────────


def _cipher(text, rows):
    mapping = {name: _entry(name, cipher) for name, cipher in rows}
    return cipher_unit_names(text, mapping)[0]


@pytest.mark.parametrize(
    "table_name",
    [
        "77 окрема танкова Тестівська бригада імені Тестових Козаків",
        "77 окрема танкова бригада імені генерал-хорунжого Тестя Тестенка",
        "77 окрема танкова бригада «Тестовий Яр»",
        "77 окрема танкова Тестівська бригада",
    ],
)
def test_honorific_only_in_table_still_ciphers(table_name):
    assert _cipher("командира роти 77 окремої танкової бригади", [(table_name, "А7777")]) == (
        "командира роти військової частини А7777"
    )


def test_full_table_name_wins_over_core_of_another_row():
    rows = [("77 окрема танкова бригада", "А0001"), ("77 окрема танкова Тестівська бригада", "А0002")]
    assert _cipher("командира 77 окремої танкової Тестівської бригади", rows) == "командира військової частини А0002"
    assert _cipher("командира 77 окремої танкової бригади", rows) == "командира військової частини А0001"


def test_ambiguous_core_is_not_guessed():
    rows = [
        ("77 окрема танкова Тестівська бригада", "А0001"),
        ("77 окрема танкова Прикладівська бригада", "А0002"),
    ]
    text = "командира 77 окремої танкової бригади"
    assert _cipher(text, rows) == text


@pytest.mark.parametrize(
    "text, expected",
    [
        (
            "командира роти 77 окремої тестової бригади імені Тестових Козаків оперативного "
            "командування «Тест» Сухопутних військ, ЗВІЛЬНИТИ",
            "командира роти військової частини А0077 оперативного командування «Тест» Сухопутних військ, ЗВІЛЬНИТИ",
        ),
        (
            "– КОМАНДИРОМ 77 ОКРЕМОЇ ТЕСТОВОЇ БРИГАДИ ІМЕНІ ТЕСТОВИХ КОЗАКІВ ОПЕРАТИВНОГО КОМАНДУВАННЯ",
            "– КОМАНДИРОМ ВІЙСЬКОВОЇ ЧАСТИНИ А0077 ОПЕРАТИВНОГО КОМАНДУВАННЯ",
        ),
        (
            "командира 77 окремої тестової бригади імені\x0bГероїв Тестівки ЗВІЛЬНИТИ з посади",
            "командира військової частини А0077 ЗВІЛЬНИТИ з посади",
        ),
        (
            "командира 77 окремої тестової бригади імені кошового отамана Тестя Тестенка, у запас",
            "командира військової частини А0077, у запас",
        ),
    ],
)
def test_unquoted_honorific_after_cipher_is_removed(text, expected):
    assert _cipher(text, [("77 окрема тестова бригада", "А0077")]) == expected


def test_honorific_not_after_cipher_is_kept():
    text = "освіта: Тестівський національний університет імені Тараса Тестенка у 2003 р."
    assert _cipher(text, [("77 окрема тестова бригада", "А0077")]) == text


def test_unnumbered_unit_with_place_adjective_is_highlighted():
    text = "офіцера окремої танкової Тестівської бригади"
    spans = generator.find_unmatched_open_unit_spans(text)
    assert [text[start:end] for start, end in spans] == ["окремої танкової Тестівської бригади"]


def test_repeat_run_with_formula_table_points_to_separate_file(tmp_path, monkeypatch):
    """Повторний запуск писав «уже є в таблиці», хоча заготовки були в окремому файлі."""
    path = _xlsx(tmp_path, [[BRIGADE, '="А"&"0077"', "77 отбр", "", "", ""]])
    monkeypatch.setattr(generator, "_append_rows_with_excel", lambda *args, **kwargs: False)
    app = generator.App.__new__(generator.App)
    logs = []
    app.log = logs.append
    app.excel_path = _Var(str(path))
    app.ADD_NEW_UNITS_TO_TABLE = True
    mapping = read_recipient_mapping(str(path))["mapping"]
    order = "НАКАЗ\n§ 1\n1. Лейтенанта ТЕСТЕНКА Теста Тестовича, офіцера 169 батальйону резерву.\n"

    app._report_table_gaps(order, mapping)
    logs.clear()
    app._report_table_gaps(order, mapping)

    assert any("нові частини.xlsx" in line for line in logs)
    assert not any("уже є в таблиці" in line for line in logs)


# ── Підпорядкування після шифру (розд. 9.5.8) ───────────────────────────────


_COMMAND_MAPPING = {
    "77 окрема тестова бригада": _entry("77 окрема тестова бригада", "А0077"),
    "оперативне командування «Тест»": _entry("оперативне командування «Тест»", "А0088"),
}


def test_command_is_ciphered_and_service_branch_is_dropped():
    text = (
        "командира роти 77 окремої тестової бригади оперативного командування «Тест» "
        "Сухопутних військ Збройних Сил України"
    )
    result, _, _ = cipher_unit_names(text, _COMMAND_MAPPING)
    assert result == "командира роти військової частини А0077 військової частини А0088"


def test_service_branch_is_dropped_in_uppercase_half():
    text = (
        "– КОМАНДИРОМ РОТИ 77 ОКРЕМОЇ ТЕСТОВОЇ БРИГАДИ ОПЕРАТИВНОГО КОМАНДУВАННЯ «ТЕСТ» "
        "СУХОПУТНИХ ВІЙСЬК ЗБРОЙНИХ СИЛ УКРАЇНИ"
    )
    result, _, _ = cipher_unit_names(text, _COMMAND_MAPPING)
    assert result == "– КОМАНДИРОМ РОТИ ВІЙСЬКОВОЇ ЧАСТИНИ А0077 ВІЙСЬКОВОЇ ЧАСТИНИ А0088"


def test_bare_armed_forces_tail_after_cipher_is_dropped():
    mapping = {"3 тестовий центр": _entry("3 тестовий центр", "А0033")}
    result, _, _ = cipher_unit_names("начальника 3 тестового центру Збройних Сил України", mapping)
    assert result == "начальника військової частини А0033"


def test_service_branch_without_cipher_before_it_is_kept():
    """Прибирається лише хвіст ПІСЛЯ шифру — решта тексту наказу недоторкана."""
    mapping = {"77 окрема тестова бригада": _entry("77 окрема тестова бригада", "А0077")}
    text = "офіцера відділу Сухопутних військ Збройних Сил України"
    result, _, _ = cipher_unit_names(text, mapping)
    assert result == text


def test_command_missing_from_table_stays_open_without_marks():
    """Рішення користувача 16.09.2026: командування лишається відкритим,
    його НЕ підсвічують і не рахують як нову частину."""
    mapping = {"77 окрема тестова бригада": _entry("77 окрема тестова бригада", "А0077")}
    text = "командира роти 77 окремої тестової бригади оперативного командування «Тест»"
    result, _, _ = cipher_unit_names(text, mapping)
    assert result == "командира роти військової частини А0077 оперативного командування «Тест»"
    assert generator.find_unmatched_open_unit_spans(result) == []
    assert generator.collect_new_unit_names(text, mapping) == []


# ── Види частин зі словника користувача (номери вигадані) ───────────────────


@pytest.mark.parametrize(
    "text, expected",
    [
        ("офіцера 12 станції фельд`єгерсько-поштового зв`язку", "12 станції"),
        ("начальника 5 окремої ремонтної майстерні засобів зв'язку", "5 окремої ремонтної майстерні"),
        ("начальника 10 командного пункту протиповітряної оборони", "10 командного пункту"),
        ("офіцера 7 командно-розвідувального пункту", "7 командно-розвідувального пункту"),
        ("офіцера 3 картографічної частини", "картографічної частини"),
        ("офіцера 44 інформаційно-телекомунікаційного вузла", "44 інформаційно-телекомунікаційного вузла"),
        ("офіцера 9 артилерійської бази боєприпасів", "9 артилерійської бази"),
        ("офіцера 1 батальйону резерву", "1 батальйону"),
    ],
)
def test_kinds_from_the_users_table_are_highlighted(text, expected):
    spans = generator.find_unmatched_open_unit_spans(text)
    assert [text[start:end] for start, end in spans] == [expected]


def test_order_reference_to_statute_part_is_not_a_unit():
    """«частина» як вид частини не годиться: ловилося б «до пункту 2 частини четвертої»."""
    text = "відповідно до пункту 2 частини четвертої статті 26 Закону України"
    assert generator.find_unmatched_open_unit_spans(text) == []


def test_only_qt_shell_adds_units_to_table():
    pytest.importorskip("PySide6")
    from nodeautomationtoolkit.generator_qt.main_window import create_qt_app_class

    assert generator.App.ADD_NEW_UNITS_TO_TABLE is False
    assert create_qt_app_class(generator).ADD_NEW_UNITS_TO_TABLE is True
