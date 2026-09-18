"""Генерація, перевірка і збірка наказу: текст за зразками додатка 53 і верстка.

Дані у фікстурах вигадані (PROJECT_RULES 1.1): справжні накази не читаються.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from nodeautomationtoolkit.order_generator.build import (  # noqa: E402
    OrderDocumentParts,
    _kind,
    build_order_document,
    content_lines,
    order_filename,
)
from nodeautomationtoolkit.order_generator.check import ERROR, WARNING, check_draft  # noqa: E402
from nodeautomationtoolkit.order_generator.compose import (  # noqa: E402
    OrderParams,
    compose_order,
    numbered_lines,
)
from nodeautomationtoolkit.order_generator.record import PLAN, PersonRecord, Position, Value  # noqa: E402

NBSP = " "


def person(
    surname="ЛІВІЦЬКИЙ",
    name="Олександр",
    patronymic="Станіславович",
    rank="полковник",
    current="начальник управління особового складу штабу",
    target="начальник управління персоналу штабу",
    shpk_from="полковник",
    shpk_to="полковник",
    vos="2905001",
    birth="1965",
    education="НАОУ (оср) у 2003 р.",
    service="09.1980",
    ipn="2649423014",
) -> PersonRecord:
    return PersonRecord(
        rank=Value(rank, PLAN),
        surname=Value(surname, PLAN),
        name=Value(name, PLAN),
        patronymic=Value(patronymic, PLAN),
        ipn=Value(ipn, PLAN),
        birth=Value(birth, PLAN),
        education=Value(education, PLAN),
        service_since=Value(service, PLAN),
        current=Position(text=Value(current, PLAN), shpk=Value(shpk_from, PLAN)),
        target=Position(
            text=Value(target, PLAN), shpk=Value(shpk_to, PLAN), vos=Value(vos, PLAN)
        ),
    )


def plain(text: str) -> str:
    return text.replace(NBSP, " ")


# ─────────────────────────────── генерація ───────────────────────────────────
def test_group_item_follows_the_sample_wording():
    draft = compose_order([person()], OrderParams(points="пункту 45", section="§ 1"))
    lines = [plain(line) for line in numbered_lines(draft)]

    assert lines[0] == "§ 1"
    assert lines[1].startswith("Відповідно до пункту 45 Положення про проходження громадянами")
    assert lines[1].endswith("ЗВІЛЬНИТИ з займаних посад і ПРИЗНАЧИТИ:")
    assert lines[3] == (
        "1. Полковника ЛІВІЦЬКОГО Олександра Станіславовича, "
        "начальника управління особового складу штабу - "
        "НАЧАЛЬНИКОМ УПРАВЛІННЯ ПЕРСОНАЛУ ШТАБУ, ВОС - 2905001."
    )
    assert lines[4] == "1965 р.н., освіта: НАОУ (оср) у 2003 р., у ЗС - із 09.1980."
    assert lines[5] == "2649423014."
    assert lines[6].startswith("Призначається на рівнозначну посаду")
    assert lines[6].endswith("з шпк «полковник» на шпк «полковник».")


def test_one_person_gets_the_single_item_wording():
    draft = compose_order([person()], OrderParams(points="пункту 45"))
    text = plain(draft.text)

    assert text.startswith("1. Відповідно до пункту 45 Положення")
    assert "ЗВІЛЬНИТИ з займаної посади і ПРИЗНАЧИТИ НАЧАЛЬНИКОМ УПРАВЛІННЯ" in text
    assert "нижчепойменованих" not in text


def test_woman_keeps_her_own_forms():
    draft = compose_order(
        [person(surname="ЗІНЧЕНКО", name="Світлана", patronymic="Петрівна", ipn="2651221542")],
        OrderParams(section="§ 1"),
    )
    assert "ЗІНЧЕНКО Світлану Петрівну" in plain(draft.text)


def test_people_are_sorted_by_surname_and_numbered_through():
    draft = compose_order(
        [person(surname="ЯНЧАК", ipn=""), person(surname="ГУРЖІЙ", ipn=""), person(surname="ЖУК", ipn="")],
        OrderParams(numbering_start=5),
    )
    assert [item.number for item in draft.items] == [5, 6, 7]
    assert [str(item.record.surname) for item in draft.items] == ["ГУРЖІЙ", "ЖУК", "ЯНЧАК"]


def test_lower_position_wording_comes_from_the_dictionary():
    record = person(shpk_from="майор", shpk_to="капітан")
    draft = compose_order([record], OrderParams(section="§ 1"))
    line = plain(draft.text).splitlines()[-1]
    assert line.startswith("Призначається на нижчу посаду")
    assert line.endswith("з шпк «майор» на шпк «капітан».")


def test_health_wording_when_the_basis_mentions_the_board():
    record = person(shpk_from="майор", shpk_to="капітан")
    record.basis = Value("висновок ВЛК від 01.09.2026", PLAN)
    draft = compose_order([record], OrderParams(section="§ 1"))
    assert "за станом здоров'я" in plain(draft.text)


# ─────────────────────────── нерозривні пробіли ──────────────────────────────
def test_non_breaking_spaces_are_applied_as_in_extracts():
    draft = compose_order([person()], OrderParams(section="§ 1"))
    text = draft.text

    assert f"ВОС{NBSP}-{NBSP}2905001" in text
    assert f"у{NBSP}ЗС" in text
    assert f"з{NBSP}шпк{NBSP}«полковник»" in text
    assert "ВОС - 2905001" not in text


# ──────────────────────────────── перевірка ──────────────────────────────────
def test_check_passes_on_a_complete_person():
    draft = compose_order([person()], OrderParams(points="пункту 45", number="525", date="17.09.2026"))
    result = check_draft(draft)
    assert result.ready
    assert not [problem for problem in result.problems if problem.level == ERROR]


def test_check_reports_the_missing_new_position():
    record = person()
    record.target = Position()
    result = check_draft(compose_order([record], OrderParams()))
    assert not result.ready
    assert any("на яку призначається" in problem.what for problem in result.errors)


def test_check_reports_a_broken_ipn():
    result = check_draft(compose_order([person(ipn="2649423011")], OrderParams()))
    assert result.ready  # це попередження, а не помилка
    assert any("РНОКПП" in problem.what for problem in result.warnings)


def test_check_reports_the_same_person_twice():
    result = check_draft(compose_order([person(), person()], OrderParams(section="§ 1")))
    assert not result.ready
    assert any("уже є в пункті" in problem.what for problem in result.errors)


def test_problem_line_speaks_plain_language():
    record = person()
    record.target = Position()
    problem = check_draft(compose_order([record], OrderParams())).errors[0]
    line = problem.line()
    assert line.startswith("✖ Пункт 1 — ЛІВІЦЬКИЙ Олександр Станіславович:")
    assert "Візьміть її з плану" in line


# ───────────────────────────────── збірка ────────────────────────────────────
def test_content_lines_keep_the_blank_line_rules():
    draft = compose_order([person(), person(surname="ЖУК", ipn="")], OrderParams(section="§ 1"))
    lines = content_lines(
        draft,
        OrderDocumentParts(
            signer_position="Командувач Сухопутних військ",
            signer_rank="генерал-лейтенант",
            signer_name="І. ПЕТРЕНКО",
            executor="Виконавець: О. КОВАЛЬ",
        ),
    )
    kinds = [_kind(line) for line in lines]

    first_item = kinds.index("item")
    second_item = kinds.index("item", first_item + 1)
    assert kinds[second_item - 1] == "blank"
    assert kinds[second_item - 2] != "blank"

    signer = kinds.index("signer")
    assert kinds[signer - 1] == "blank" and kinds[signer - 2] == "blank"
    assert kinds[signer - 3] != "blank"
    assert lines[-1].strip()


def test_built_document_follows_the_layout_rules(tmp_path):
    from docx import Document
    from docx.shared import Pt

    draft = compose_order([person(), person(surname="ЖУК", ipn="")], OrderParams(section="§ 1"))
    path = build_order_document(
        draft,
        tmp_path / "order.docx",
        parts=OrderDocumentParts(
            signer_position="Командувач Сухопутних військ",
            signer_rank="генерал-лейтенант",
            signer_name="І. ПЕТРЕНКО",
            executor="Виконавець: О. КОВАЛЬ",
        ),
    )
    document = Document(str(path))
    section = document.sections[0]
    # Word зберігає поля у твіпах, тому назад вони приходять із похибкою округлення.
    assert round(section.left_margin.cm, 2) == 2.0
    assert round(section.right_margin.cm, 2) == 1.0

    paragraphs = [p for p in document.paragraphs]
    body = [p for p in paragraphs if p.text.strip()]
    for paragraph in body:
        for run in paragraph.runs:
            assert run.font.name == "Times New Roman"

    items = [p for p in paragraphs if _kind(p.text) == "item"]
    assert items, "у документі немає пунктів"
    for paragraph in items:
        assert paragraph.paragraph_format.keep_together is True
        assert round(paragraph.paragraph_format.first_line_indent.cm, 2) == 1.25

    executor = paragraphs[-1]
    assert executor.runs and executor.runs[0].font.size == Pt(8)
    assert paragraphs[-1].text.strip(), "у кінці документа не повинно бути порожніх абзаців"


def test_heading_is_kept_with_the_first_item(tmp_path):
    from docx import Document

    draft = compose_order([person(), person(surname="ЖУК", ipn="")], OrderParams(section="§ 1"))
    path = build_order_document(draft, tmp_path / "order.docx")
    paragraphs = Document(str(path)).paragraphs
    heading = next(p for p in paragraphs if p.text.strip().endswith("ПРИЗНАЧИТИ:"))
    assert heading.paragraph_format.keep_with_next is True


def test_template_content_tag_is_replaced(tmp_path):
    from docx import Document

    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("НАКАЗ № {{номер_наказу}} від {{дата_наказу}}")
    document.add_paragraph("{{зміст}}")
    document.add_paragraph("Командувач Сухопутних військ")
    document.save(str(template))

    draft = compose_order([person()], OrderParams(number="525", date="17.09.2026"))
    path = build_order_document(draft, tmp_path / "out.docx", template)
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)

    assert "{{зміст}}" not in text
    assert "{{номер_наказу}}" not in text
    assert "НАКАЗ № 525 від 17.09.2026" in text
    assert "ЛІВІЦЬКОГО Олександра Станіславовича" in text
    assert text.strip().endswith("Командувач Сухопутних військ")


def test_order_filename_uses_the_requisites():
    draft = compose_order([person()], OrderParams(number="525", date="17.09.2026"))
    assert order_filename(draft) == "Наказ № 525 від 17.09.2026.docx"


@pytest.mark.parametrize(
    "line, kind",
    [
        ("§ 1", "heading"),
        ("Відповідно до пункту 45 … ПРИЗНАЧИТИ:", "heading"),
        ("1. Полковника ЛІВІЦЬКОГО …", "item"),
        ("1965 р.н., освіта: …", "continuation"),
        ("Командувач Сухопутних військ", "signer"),
        ("", "blank"),
    ],
)
def test_paragraph_kinds(line, kind):
    assert _kind(line) == kind
