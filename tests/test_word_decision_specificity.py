from docx import Document

from nodeautomationtoolkit.builtin_nodes.word import generate_decision_order_docx


def test_word_decision_node_prefers_specific_column_a_name_and_keeps_anaphora(tmp_path):
    source_path = tmp_path / "source.docx"
    output_path = tmp_path / "result.docx"
    source = Document()
    source.add_paragraph("Службова шапка")
    source.add_paragraph(
        "1. Офіцера 908 окремого центру підготовки призначити "
        "начальником служби цього самого центру."
    )
    source.save(source_path)

    mapping = {
        "окремий центр підготовки": {
            "cipher": "А0002",
            "abbreviation": "оцп",
        },
        "908 центр підготовки": {
            "cipher": "А0908",
            "abbreviation": "908 цп",
        },
        "центр": {
            "cipher": "А0003",
            "abbreviation": "центр",
        },
    }

    generate_decision_order_docx(
        path=str(source_path),
        output_path=str(output_path),
        mapping=mapping,
        new_header="Службова шапка",
    )

    result = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)
    assert "А0908" in result
    assert "А0002" not in result
    assert "А0003" not in result
    assert "цієї самої військової частини" in result
