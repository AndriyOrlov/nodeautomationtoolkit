"""Наказ про присвоєння військових звань: давальний відмінок, групи за званням.

Формулювання — зі зразків додатка 53; дані вигадані (PROJECT_RULES 1.1).
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from nodeautomationtoolkit.order_generator.check import check_draft  # noqa: E402
from nodeautomationtoolkit.order_generator.compose import (  # noqa: E402
    RANK,
    OrderParams,
    compose_order,
)
from nodeautomationtoolkit.order_generator.names import (  # noqa: E402
    dative,
    split_full_name,
)
from nodeautomationtoolkit.order_generator.record import (  # noqa: E402
    MANUAL,
    PersonRecord,
    Position,
    Value,
)
from nodeautomationtoolkit.personnel import declension  # noqa: E402

NBSP = " "


def person(
    surname="НАЗАРЕНКО",
    name="Олександр",
    patronymic="Васильович",
    rank="капітан",
    position="старший офіцер відділу комплектування офіцерами",
    new_rank="майор",
    seniority="11 років",
    ipn="2859220555",
    birth="1978",
    since="",
    note="",
    shpk="",
) -> PersonRecord:
    return PersonRecord(
        rank=Value(rank, MANUAL),
        surname=Value(surname, MANUAL),
        name=Value(name, MANUAL),
        patronymic=Value(patronymic, MANUAL),
        ipn=Value(ipn, MANUAL),
        birth=Value(birth, MANUAL),
        current=Position(text=Value(position, MANUAL)),
        target=Position(shpk=Value(shpk, MANUAL)),
        new_rank=Value(new_rank, MANUAL),
        rank_seniority=Value(seniority, MANUAL),
        rank_since=Value(since, MANUAL),
        rank_note=Value(note, MANUAL),
    )


def params(**extra) -> OrderParams:
    values = {
        "action": RANK,
        "points": "пунктів 45, 50",
        "kind": "особам офіцерського складу",
        "unit": "Збройних Сил України",
    }
    values.update(extra)
    return OrderParams(**values)


def plain(text: str) -> str:
    return text.replace(NBSP, " ")


# ─────────────────────────────── відмінок ────────────────────────────────────
def test_dative_full_names():
    assert dative(split_full_name("НАЗАРЕНКО Олександр Васильович")) == (
        "НАЗАРЕНКУ Олександру Васильовичу"
    )
    assert dative(split_full_name("ДУБ Романна Іванівна")) == "ДУБ Романні Іванівні"
    assert dative(split_full_name("КИРИЧЕНКО Леся Петрівна")) == "КИРИЧЕНКО Лесі Петрівні"
    assert dative(split_full_name("ЛІВІЦЬКИЙ Олександр Станіславович")) == (
        "ЛІВІЦЬКОМУ Олександру Станіславовичу"
    )


def test_dictionaries_know_the_dative():
    assert declension.decline_rank("капітан", "Д").text == "капітану"
    assert declension.decline_rank("молодший сержант", "Д").text == "молодшому сержанту"
    position = declension.decline_position("старший офіцер відділу", "Д")
    assert position.text == "старшому офіцеру відділу" and position.found


# ─────────────────────────────── генерація ───────────────────────────────────
def test_rank_item_follows_the_sample_wording():
    draft = compose_order([person()], params())
    lines = [plain(line) for line in draft.text.splitlines()]

    assert lines[0].endswith("ПРИСВОЇТИ чергові військові звання:")
    assert "нижчепойменованим особам офіцерського складу" in lines[0]
    assert lines[2] == "«МАЙОР»"
    assert lines[4] == (
        "1. Капітану НАЗАРЕНКУ Олександру Васильовичу, старшому офіцеру відділу "
        "комплектування офіцерами."
    )
    assert lines[5] == "1978 р.н., вислуга у званні - 11 років, 2859220555."


def test_seniority_start_line():
    draft = compose_order([person(since="04.12.2026")], params())
    assert plain(draft.text).endswith(
        "Строк перебування у військовому званні рахувати з 04.12.2026."
    )


def test_ahead_of_time_note_and_shpk():
    draft = compose_order(
        [person(seniority="", note="достроково на 6 місяців", shpk="полковник")], params()
    )
    assert "1978 р.н., достроково на 6 місяців, 2859220555, шпк «полковник»." in plain(draft.text)


def test_people_are_grouped_by_the_new_rank():
    draft = compose_order(
        [
            person(surname="ЗОРЕНКО", new_rank="полковник", ipn=""),
            person(surname="ШАПРАН", new_rank="майор", ipn=""),
            person(surname="НАЗАРЕНКО", new_rank="майор", ipn=""),
        ],
        params(),
    )
    assert [item.number for item in draft.items] == [1, 2, 3]
    assert [item.subheading for item in draft.items] == ["«МАЙОР»", "«МАЙОР»", "«ПОЛКОВНИК»"]
    assert [str(item.record.surname) for item in draft.items] == [
        "НАЗАРЕНКО",
        "ШАПРАН",
        "ЗОРЕНКО",
    ]
    assert plain(draft.text).count("«МАЙОР»") == 1


def test_rank_order_has_no_appointment_or_dismissal_wording():
    text = plain(compose_order([person()], params()).text)
    assert "ПРИЗНАЧИТИ" not in text and "ЗВІЛЬНИТИ" not in text


# ──────────────────────────────── перевірка ──────────────────────────────────
def test_check_passes_on_a_complete_rank_item():
    draft = compose_order([person()], params(number="531", date="18.09.2026"))
    result = check_draft(draft)
    assert result.ready, [problem.line() for problem in result.errors]


def test_check_wants_the_new_rank():
    result = check_draft(compose_order([person(new_rank="")], params()))
    assert not result.ready
    assert any("яке звання присвоюється" in problem.what for problem in result.errors)


def test_check_does_not_ask_for_a_new_position():
    result = check_draft(compose_order([person()], params()))
    assert not any("на яку призначається" in problem.what for problem in result.problems)


def test_missing_position_is_only_a_warning_here():
    result = check_draft(compose_order([person(position="")], params()))
    assert result.ready
    assert any("посаду, яку обіймає" in problem.what for problem in result.warnings)


def test_check_warns_about_the_missing_seniority():
    result = check_draft(compose_order([person(seniority="")], params()))
    assert result.ready
    assert any("вислуги у званні" in problem.what for problem in result.warnings)


# ───────────────────────────────── збірка ────────────────────────────────────
def test_rank_subheading_is_kept_with_the_first_item(tmp_path):
    from docx import Document

    from nodeautomationtoolkit.order_generator.build import build_order_document

    draft = compose_order(
        [person(), person(surname="ЗОРЕНКО", new_rank="полковник", ipn="")], params()
    )
    path = build_order_document(draft, tmp_path / "ranks.docx")
    paragraphs = Document(str(path)).paragraphs
    heading = next(p for p in paragraphs if p.text.strip() == "«МАЙОР»")
    assert heading.paragraph_format.keep_with_next is True


def test_rank_subheading_is_centred(tmp_path):
    """У зразку додатка 53 звання-підзаголовок стоїть по центру аркуша."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from nodeautomationtoolkit.order_generator.build import build_order_document

    draft = compose_order([person()], params())
    path = build_order_document(draft, tmp_path / "ranks.docx")
    paragraphs = Document(str(path)).paragraphs
    heading = next(p for p in paragraphs if p.text.strip() == "«МАЙОР»")
    assert heading.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert round(heading.paragraph_format.left_indent.cm, 2) == 0.0

    biography = next(p for p in paragraphs if "вислуга" in plain(p.text))
    assert round(biography.paragraph_format.left_indent.cm, 1) == 8.0
