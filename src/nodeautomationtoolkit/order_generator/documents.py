"""М'який розбір будь-якого кадрового документа.

Кидають файл — програма читає текст (.docx, .doc, .rtf, .txt, .xlsx), визначає
вид документа й витягує те, що там є: ПІБ, звання, РНОКПП, посади, шпк, ВОС,
тарифний розряд, дати, згоду на призначення, підстави. Нічого обов'язкового
немає: чого немає в документі, добирається з наказів через індекс, а що є —
звіряється з ними (`resolve.py`).

Види документів — за Інструкцією про організацію виконання Положення про
проходження військової служби (наказ МО України від 10.04.2009 № 170):
подання (додатки 1, 3), список до присвоєння звання (додаток 4), перелік посад
(додаток 6), службова характеристика (додаток 7), аркуш вивчення особистих
якостей (додаток 8), картка професійного відбору (додаток 9), довідка про
проходження служби (додаток 10), атестація (додаток 11), аркуш бесіди
(додаток 12), оцінна картка (додатки 13, 14), план переміщення (додаток 16),
план звільнення (додаток 21), рапорт (заява) на звільнення (додаток 22),
акт перевірки сімейного стану (додаток 23), довідка (додаток 24).
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from ..order_index.extractor import normalize_space, read_docx_paragraphs
from ..order_index.items import _PERSON_RE  # «ПРІЗВИЩЕ Ім'я По батькові,»

_UP = "А-ЯІЇЄҐ"
_LO = "а-яіїєґ"

#: Вид документа → слова, за якими він упізнається (перше збіжне й перемагає).
DOCUMENT_KINDS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("план переміщення", ("план переміщення",)),
    ("план звільнення", ("план звільнення",)),
    ("наказ", ("наказ по особовому складу", "наказую:", "звільнити з займаних посад", "§")),
    ("витяг", ("витяг з наказу", "витяг із наказу")),
    ("подання до звільнення", ("подання", "до звільнення")),
    ("подання до присвоєння звання", ("подання", "присвоєння військового звання")),
    ("подання", ("подання",)),
    ("рапорт", ("рапорт", "заява", "прошу", "бажаю звільнитися")),
    ("аркуш бесіди", ("аркуш бесіди",)),
    ("аркуш вивчення", ("аркуш вивчення особистих якостей",)),
    ("картка професійного відбору", ("картка професійного відбору",)),
    ("службова характеристика", ("службова характеристика",)),
    ("атестація", ("атестація",)),
    ("довідка", ("довідка",)),
    ("акт", ("акт перевірки",)),
    ("список", ("список військовослужбовців",)),
    ("перелік посад", ("перелік посад",)),
)

_IPN_RE = re.compile(r"(?<!\d)(\d{10})(?!\d)")
_BIRTH_RE = re.compile(r"(?<!\d)(\d{2}\.\d{2}\.(?:19|20)\d{2})(?!\d)|(?<!\d)((?:19|20)\d{2})\s*р\.\s*н\.")
_SERVICE_RE = re.compile(r"у\s+ЗС\w*\s*[-–]?\s*(?:із|з)\s+(\d{2}\.\d{4}|\d{4})", re.IGNORECASE)
_EDUCATION_RE = re.compile(r"освіт[аи][:\s-]+(.{5,160}?)(?:;|\.\s|$)", re.IGNORECASE)
_SHPK_RE = re.compile(r"шпк\s*[-–:]?\s*[«\"“„]?([^»\"”,;)]+)[»\"”]?", re.IGNORECASE)
_VOS_RE = re.compile(r"ВОС\s*[-–]?\s*(\d{6,8}[А-ЯA-Z]?)", re.IGNORECASE)
_TARIFF_RE = re.compile(r"(\d{1,2})\s*(?:т\.?\s?р\.?|тарифн\w*\s+розряд\w*)", re.IGNORECASE)
_ORDER_REF_RE = re.compile(r"наказ\w*[^.;]{0,80}?від\s+(\d{2}\.\d{2}\.\d{4})\s*№\s*([\w\-/]+)", re.IGNORECASE)
_REPORT_REF_RE = re.compile(r"(рапорт|заяв\w+)[^.;]{0,40}?від\s+(\d{2}\.\d{2}\.\d{4})", re.IGNORECASE)
_LAW_REF_RE = re.compile(
    r"(?:під)?пункт\w*\s*[«\"]?([\w\-]+)[»\"]?\s*(?:частини\s+\w+\s+)?статті\s+(\d+)[^.;]{0,60}", re.IGNORECASE
)
_CONSENT_RE = re.compile(
    r"(з[гґ]од(?:ен|на|ий|а)\b[^.;]{0,60}|бажаю[^.;]{0,60}|прошу\s+(?:призначити|перемістити|звільнити)[^.;]{0,80})",
    re.IGNORECASE,
)
_DISCHARGE_RE = re.compile(
    r"(звільн\w+[^.;]{0,80}|виключ\w+\s+зі\s+списків[^.;]{0,60})", re.IGNORECASE
)
_APPOINT_RE = re.compile(r"(призначи\w+|перемісти\w+|признач\w+)[^.;]{0,120}", re.IGNORECASE)


@dataclass
class Fact:
    """Знайдене значення: що саме, звідки в тексті і як упевнено."""

    field: str
    value: str
    context: str = ""
    confidence: str = "м'яко"  # «точно» — однозначна ознака, «м'яко» — здогад


@dataclass
class PersonMention:
    rank: str = ""
    surname: str = ""
    name: str = ""
    patronymic: str = ""
    ipn: str = ""
    context: str = ""

    @property
    def full_name(self) -> str:
        return " ".join(part for part in (self.surname, self.name, self.patronymic) if part)


@dataclass
class DocumentFacts:
    path: str = ""
    kind: str = "невідомо"
    text: str = ""
    people: list[PersonMention] = field(default_factory=list)
    facts: list[Fact] = field(default_factory=list)
    positions: list[str] = field(default_factory=list)

    def first(self, field_name: str) -> str:
        return next((fact.value for fact in self.facts if fact.field == field_name), "")

    def all(self, field_name: str) -> list[str]:
        return [fact.value for fact in self.facts if fact.field == field_name]


def read_text(path: str | Path, word_reader=None) -> str:
    """Текст будь-якого кадрового файла; .doc/.rtf — через Word (`word_reader`)."""
    path = Path(path)
    suffix = path.suffix.casefold()
    if suffix in {".txt", ".csv", ".md"}:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", "replace")
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [normalize_space(str(value)) for value in row if value not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    if suffix in {".docx", ".docm"} or zipfile.is_zipfile(path):
        return "\n".join(read_docx_paragraphs(path))
    if word_reader is not None:
        return "\n".join(word_reader.read_paragraphs(path))
    from ..order_index.word_reader import WordTextReader

    with WordTextReader() as reader:
        return "\n".join(reader.read_paragraphs(path))


def detect_kind(text: str) -> str:
    """Вид документа за його текстом; «невідомо», якщо ознак немає."""
    head = " ".join(text[:4000].split()).casefold()
    for kind, markers in DOCUMENT_KINDS:
        if all(marker.casefold() in head for marker in markers):
            return kind
    return "невідомо"


def _people(text: str) -> list[PersonMention]:
    people: list[PersonMention] = []
    seen: set[str] = set()
    for match in _PERSON_RE.finditer(text):
        words = match.group(0).strip(" ,").split()
        if len(words) < 3:
            continue
        before = text[max(0, match.start() - 120):match.start()]
        after = text[match.end():match.end() + 160]
        person = PersonMention(
            rank=normalize_space(before.split(",")[-1])[-40:].strip(),
            surname=words[0],
            name=words[1],
            patronymic=words[2],
            context=normalize_space(before[-60:] + match.group(0) + after)[:220],
        )
        ipn = _IPN_RE.search(after)
        if ipn:
            person.ipn = ipn.group(1)
        key = person.full_name.casefold()
        if key not in seen:
            seen.add(key)
            people.append(person)
    return people


def _add(facts: list[Fact], field_name: str, value: str, context: str, confidence: str = "м'яко") -> None:
    value = normalize_space(value)
    if value and not any(fact.field == field_name and fact.value == value for fact in facts):
        facts.append(Fact(field_name, value, normalize_space(context)[:220], confidence))


def extract_facts(text: str, path: str = "", dictionary=None) -> DocumentFacts:
    """Усе, що вдалося дістати з документа; нічого обов'язкового немає."""
    text = text.replace("\xa0", " ")
    result = DocumentFacts(path=path, kind=detect_kind(text), text=text)
    result.people = _people(text)

    for match in _IPN_RE.finditer(text):
        _add(result.facts, "ipn", match.group(1), _around(text, match), "точно")
    for match in _BIRTH_RE.finditer(text):
        _add(result.facts, "birth", match.group(1) or match.group(2), _around(text, match))
    for pattern, field_name, confidence in (
        (_SERVICE_RE, "service_since", "точно"),
        (_EDUCATION_RE, "education", "точно"),
        (_SHPK_RE, "shpk", "точно"),
        (_VOS_RE, "vos", "точно"),
        (_TARIFF_RE, "tariff", "точно"),
        (_CONSENT_RE, "consent", "м'яко"),
        (_DISCHARGE_RE, "discharge", "м'яко"),
        (_APPOINT_RE, "appointment", "м'яко"),
    ):
        for match in pattern.finditer(text):
            _add(result.facts, field_name, match.group(1) if match.groups() else match.group(0),
                 _around(text, match), confidence)
    for match in _ORDER_REF_RE.finditer(text):
        _add(result.facts, "order_reference", f"наказ від {match.group(1)} № {match.group(2)}",
             _around(text, match), "точно")
    for match in _REPORT_REF_RE.finditer(text):
        _add(result.facts, "basis", f"{match.group(1)} від {match.group(2)}", _around(text, match), "точно")
    for match in _LAW_REF_RE.finditer(text):
        _add(result.facts, "law_reference", match.group(0), _around(text, match), "точно")

    result.positions = _positions(text, dictionary)
    for position in result.positions:
        _add(result.facts, "position", position, position)
    return result


def _around(text: str, match: re.Match, width: int = 90) -> str:
    return text[max(0, match.start() - width):match.end() + width]


def _positions(text: str, dictionary) -> list[str]:
    """Посади, які впізнав словник, у порядку появи (без повторів)."""
    if dictionary is None:
        from ..order_index.position_dictionary import load_position_dictionary

        dictionary = load_position_dictionary()
    found: list[str] = []
    for piece in re.split(r"[\n;]|(?<=[.,])\s", text):
        piece = piece.strip(" ,.;:-–—")
        if len(piece) < 5:
            continue
        match = dictionary.match_at_start(piece)
        if match and match.nominative not in found:
            found.append(match.nominative)
    return found


def read_document(path: str | Path, word_reader=None, dictionary=None) -> DocumentFacts:
    """Прочитати файл і витягти з нього все, що вдасться."""
    text = read_text(path, word_reader)
    return extract_facts(text, str(path), dictionary)
