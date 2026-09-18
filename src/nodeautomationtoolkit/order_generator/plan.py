"""Розбір плану переміщення військовослужбовців (додаток 16 до Інструкції, наказ МО № 170).

Графи таблиці:
1 — № з/п;
2 — посада, що підлягає комплектуванню (шпк, ВОС, тарифний розряд, з якої дати вакантна
    або хто та коли її вивільняє);
3 — кандидат: звання (дата присвоєння), ПРІЗВИЩЕ Ім'я По батькові, РНОКПП (особистий номер),
    займана посада, з якого часу, [колишня посада], (шпк, ВОС, тарифний розряд);
4 — дата народження, освіта (рік закінчення), [у ЗС - із …], [бойовий досвід];
5 — висновок і бали щорічного оцінювання;
6 — підстава включення до плану; 7 — відмітка про реалізацію.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

_UP = "А-ЯІЇЄҐ"
_PARAMS_RE = re.compile(r"\(\s*(шпк[^()]*?)\)", re.IGNORECASE)
_SHPK_RE = re.compile(r"шпк\s*[-–:]?\s*[«\"“„]([^»\"”]+)[»\"”]", re.IGNORECASE)
_VOS_RE = re.compile(r"ВОС\s*[-–]?\s*(\d{6,8}[А-ЯA-Z]?)", re.IGNORECASE)
_TARIFF_RE = re.compile(r"(\d{1,2})\s*(?:т\.?\s?р\.?|грн)?\s*$")
_DATE = r"\d{2}\.\d{2}\.\d{4}"
_CANDIDATE_RE = re.compile(
    rf"^(?P<rank>[^()]+?)\s*\(\s*(?P<rank_date>{_DATE})\s*\)\s*"
    rf"(?P<surname>[{_UP}][{_UP}'’ʼ-]+)\s+(?P<name>[{_UP}][^\s,]+)\s+(?P<patronymic>[{_UP}][^\s,]+)\s*,\s*"
    rf"(?:(?P<ipn>\d{{10}}|[{_UP}A-Z]{{1,2}}[-\s]?\d{{6,9}})\s*,\s*)?(?P<rest>.*)$",
    re.DOTALL,
)
_SINCE_RE = re.compile(rf"\s+з\s+({_DATE})")
_FORMER_RE = re.compile(r",\s*колишн(?:ій|ьої|ього|я)\s+(.+)$", re.IGNORECASE | re.DOTALL)
_BIRTH_RE = re.compile(rf"^\s*({_DATE}|\d{{4}}\s*р\.?\s*н\.?)\s*,?\s*")
_SERVICE_RE = re.compile(r",?\s*у\s+ЗС\s*[-–]?\s*(?:із|з)\s+(\d{2}\.\d{4})\.?", re.IGNORECASE)
_EXPERIENCE_RE = re.compile(r",?\s*(?:бойовий\s+досвід|учасник\s+бойових\s+дій)[^,.]*[.,]?", re.IGNORECASE)


@dataclass
class PositionParams:
    shpk: str = ""
    vos: str = ""
    tariff: str = ""


@dataclass
class Vacancy:
    position: str = ""  # у називному, разом із частиною
    params: PositionParams = field(default_factory=PositionParams)
    note: str = ""  # «вакантна з …» або хто й коли вивільняє


@dataclass
class Candidate:
    rank: str = ""
    rank_date: str = ""
    surname: str = ""
    name: str = ""
    patronymic: str = ""
    ipn: str = ""
    position: str = ""  # займана посада в називному, разом із частиною
    since: str = ""
    former_position: str = ""
    params: PositionParams = field(default_factory=PositionParams)


@dataclass
class Biography:
    birth: str = ""  # «11.07.1983» або «1983 р.н.»
    education: str = ""
    service_since: str = ""  # «08.1996»
    experience: str = ""

    @property
    def birth_year(self) -> str:
        match = re.search(r"\d{4}", self.birth)
        return match.group(0) if match else ""


@dataclass
class PlanEntry:
    number: str
    vacancy: Vacancy
    candidate: Candidate
    biography: Biography
    evaluation: str
    basis: str
    raw: tuple[str, ...]
    source: str = ""
    row: int = 0
    problems: list[str] = field(default_factory=list)


@dataclass
class Plan:
    path: Path
    title: str
    approved: str  # рядок затвердження («ЗАТВЕРДЖУЮ … «11» грудня 2021 року»)
    entries: list[PlanEntry]
    problems: list[str] = field(default_factory=list)


def _norm(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\xa0", " ").replace("’", "'").replace("ʼ", "'")
    return " ".join(text.split())


def parse_params(text: str) -> tuple[PositionParams, str]:
    """Дужки «(шпк «майор», ВОС-0304003, 25)» → параметри і текст без них."""
    match = None
    for match in _PARAMS_RE.finditer(text):
        pass  # беремо останні дужки з шпк
    if not match:
        return PositionParams(), text
    inner = match.group(1)
    shpk = _SHPK_RE.search(inner)
    vos = _VOS_RE.search(inner)
    rest = inner
    for found in (shpk, vos):
        if found:
            rest = rest.replace(found.group(0), " ")
    tariff = _TARIFF_RE.search(rest.strip(" ,;"))
    params = PositionParams(
        shpk=shpk.group(1).strip() if shpk else "",
        vos=vos.group(1) if vos else "",
        tariff=tariff.group(1) if tariff else "",
    )
    return params, (text[: match.start()] + text[match.end():]).strip(" ,")


def parse_vacancy(text: str) -> Vacancy:
    text = _norm(text)
    params, rest = parse_params(text)
    match = re.search(r"\(\s*шпк", text, re.IGNORECASE)
    if match:
        position = text[: match.start()].strip(" ,")
        note = text[match.start():]
        note = _PARAMS_RE.sub("", note, count=1).strip(" ,")
    else:
        position, _, note = rest.partition(", ")
    return Vacancy(position=position, params=params, note=note)


def parse_candidate(text: str) -> tuple[Candidate, list[str]]:
    text = _norm(text)
    problems: list[str] = []
    params, rest_text = parse_params(text)
    match = _CANDIDATE_RE.match(rest_text)
    if not match:
        return Candidate(params=params), ["не розпізнано звання (дата) ПРІЗВИЩЕ Ім'я По батькові"]
    candidate = Candidate(
        rank=match.group("rank").strip(" ,"),
        rank_date=match.group("rank_date"),
        surname=match.group("surname"),
        name=match.group("name"),
        patronymic=match.group("patronymic"),
        ipn=match.group("ipn") or "",
        params=params,
    )
    rest = match.group("rest").strip(" ,")
    former = _FORMER_RE.search(rest)
    if former:
        candidate.former_position = former.group(1).strip(" ,.")
        rest = rest[: former.start()]
    since = _SINCE_RE.search(rest)
    if since:
        candidate.since = since.group(1)
        rest = rest[: since.start()] + rest[since.end():]
    candidate.position = rest.strip(" ,.")
    if not candidate.ipn:
        problems.append("немає РНОКПП")
    if not params.shpk:
        problems.append("немає шпк займаної посади")
    return candidate, problems


def parse_biography(text: str) -> Biography:
    text = _norm(text)
    bio = Biography()
    birth = _BIRTH_RE.match(text)
    if birth:
        bio.birth = birth.group(1)
        text = text[birth.end():]
    service = _SERVICE_RE.search(text)
    if service:
        bio.service_since = service.group(1)
        text = text[: service.start()] + text[service.end():]
    experience = _EXPERIENCE_RE.search(text)
    if experience:
        bio.experience = experience.group(0).strip(" ,.")
        text = text[: experience.start()] + text[experience.end():]
    bio.education = text.strip(" ,")
    if bio.education and not bio.education.endswith("."):
        bio.education += "."
    return bio


def _header_row(sheet) -> tuple[int, int] | None:
    """(номер рядка, стовпець «№») шапки таблиці."""
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 40)):
        for cell in row:
            if "найменування посади" in _norm(cell.value).casefold():
                return cell.row, max(cell.column - 1, 1)
    return None


def read_plan(path: str | Path) -> Plan:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook.worksheets[0]
    header = _header_row(sheet)
    title_lines = []
    for row in sheet.iter_rows(min_row=1, max_row=(header[0] - 1) if header else 15, values_only=True):
        line = " ".join(_norm(value) for value in row if _norm(value))
        if line:
            title_lines.append(line)
    title = next((line for line in title_lines if line.upper().startswith("ПЛАН")), "")
    approved = " ".join(title_lines[: title_lines.index(title)]) if title in title_lines else ""
    plan = Plan(path=path, title=title, approved=approved, entries=[])
    if header is None:
        plan.problems.append("не знайдено шапку таблиці («Найменування посади…»)")
        return plan
    header_row, first_column = header
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row):
        values = [_norm(row[first_column - 1 + offset].value) if first_column - 1 + offset < len(row) else ""
                  for offset in range(7)]
        number, vacancy_text, candidate_text = values[0], values[1], values[2]
        if not re.fullmatch(r"\d{1,4}\.?", number) or not (vacancy_text and candidate_text):
            continue
        candidate, problems = parse_candidate(candidate_text)
        entry = PlanEntry(
            number=number.rstrip("."),
            vacancy=parse_vacancy(vacancy_text),
            candidate=candidate,
            biography=parse_biography(values[3]),
            evaluation=values[4],
            basis=values[5],
            raw=tuple(values),
            source=path.name,
            row=row[0].row,
            problems=problems,
        )
        if not entry.vacancy.params.shpk:
            entry.problems.append("немає шпк посади, що комплектується")
        plan.entries.append(entry)
    return plan
