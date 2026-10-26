"""Excel-звіти розрахунку розсилки й контрольні файли.

Винесено з `generate_extracts.py` без змін логіки (26.10.2026).
"""
from __future__ import annotations

import os
import hashlib
import tempfile
import time
import openpyxl
from nodeautomationtoolkit.builtin_nodes.recipient_mapping import (
    _format_item_numbers_range,
)

from .naming import sanitize_filename


def _save_table_to_excel(filepath: str, headers: list[str], rows: list):
    """Зберігає таблицю в Excel (.xlsx) за допомогою openpyxl."""
    os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        if isinstance(row, dict):
            ws.append([row.get(h, "") for h in headers])
        else:
            ws.append(list(row))
    try:
        wb.save(filepath)
    except PermissionError:
        raise PermissionError(
            f"Файл '{os.path.basename(filepath)}' зайнятий іншою програмою. "
            f"Будь ласка, закрийте його в Excel і спробуйте знову."
        )


def select_extracts_for_scope(map_result: dict, scope: str = "all") -> tuple[dict, dict]:
    """Повертає адресні й управлінські витяги для вибраного режиму запуску.

    Маршрутизатор завжди аналізує наказ повністю. Розділення робимо вже після
    аналізу, щоб окрема кнопка управлінських витягів не створювала другий набір
    правил розпізнавання й не могла розійтися зі звичайною генерацією.
    """
    units = map_result.get("unit_paragraphs", {})
    management = map_result.get("management_paragraphs", {})
    if scope == "general":
        return units, {}
    if scope == "management":
        return {}, management
    if scope == "all":
        return units, management
    raise ValueError(f"Невідомий режим генерації витягів: {scope}")


MANAGEMENT_RESULT_LABEL = "Управління · окремий витяг"


def management_calculation_rows(map_result: dict) -> list[tuple[str, str, int]]:
    """Інформаційні рядки управління для таблиці результатів у програмі.

    Вони навмисно не додаються до XLSX розрахунку розсилки: це не адресати й
    не відправлення. Зелений рядок у UI лише показує, що пункт розпізнано та
    для нього можна сформувати окремий витяг до управління.
    """
    rows = []
    for data in map_result.get("management_paragraphs", {}).values():
        items = data.get("items", []) if isinstance(data, dict) else []
        labels = [item.get("label", "") for item in items if isinstance(item, dict)]
        rows.append((MANAGEMENT_RESULT_LABEL, _format_item_numbers_range(labels), len(items)))
    return rows


def _saved_xlsx_rows(path: str) -> list[tuple]:
    """Читає вже створений контрольний XLSX без будь-якого перерахунку."""
    if not os.path.isfile(path):
        return []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [
            tuple("" if value is None else value for value in row)
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(value not in (None, "") for value in row)
        ]
    finally:
        workbook.close()


#: Контрольні файли, старші за цей строк, видаляються з %TEMP% при першому
#: зверненні за сеанс — щоб тимчасова тека не росла роками.
CONTROL_REPORTS_MAX_AGE_DAYS = 30


_control_reports_cleaned = False


def cleanup_control_reports(root: str, max_age_days: int = CONTROL_REPORTS_MAX_AGE_DAYS) -> int:
    """Видаляє старі контрольні файли й порожні підтеки; повертає кількість файлів."""
    cutoff = time.time() - max_age_days * 86400
    removed = 0
    if not os.path.isdir(root):
        return 0
    for name in os.listdir(root):
        folder = os.path.join(root, name)
        if not os.path.isdir(folder):
            continue
        for file_name in os.listdir(folder):
            path = os.path.join(folder, file_name)
            try:
                if os.path.isfile(path) and os.path.getmtime(path) < cutoff:
                    os.remove(path)
                    removed += 1
            except OSError:
                continue  # файл відкритий у Excel — спробуємо наступного разу
        try:
            os.rmdir(folder)  # лише якщо тека спорожніла
        except OSError:
            pass
    return removed


def control_reports_folder(output_folder: str) -> str:
    """Тимчасова тека для службових контрольних файлів одного результату.

    У теці результату користувачу потрібен лише «Розрахунок розсилки»
    (вимога 26.10.2026). Контроль маршрутизації, пропущених пунктів і складу
    витягів програма читає сама (вкладки вікна), тому вони лежать у %TEMP%.
    Окрема підтека на кожну теку результату: однакові назви наказів із різних
    тек не перезаписують одна одну.
    """
    global _control_reports_cleaned
    root = os.path.join(tempfile.gettempdir(), "GeneratorVytyagiv_control")
    if not _control_reports_cleaned:
        _control_reports_cleaned = True
        cleanup_control_reports(root)
    key = os.path.normcase(os.path.abspath(output_folder or "."))
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:12]
    folder = os.path.join(root, digest)
    os.makedirs(folder, exist_ok=True)
    return folder


def control_report_path(output_folder: str, prefix: str, order_base: str) -> str:
    """Шлях контрольного файлу `<prefix>_<наказ>.xlsx` у тимчасовій теці."""
    return os.path.join(control_reports_folder(output_folder), f"{prefix}_{order_base}.xlsx")


def load_saved_analysis_reports(order_path: str, output_folder: str) -> dict:
    """Завантажує збережені розсилку, пропуски й маршрутизацію для наказу."""
    order_base = sanitize_filename(os.path.splitext(os.path.basename(order_path))[0])

    def control(prefix: str) -> str:
        # Нове місце — %TEMP%; файли старих запусків лишились у теці результату.
        path = control_report_path(output_folder, prefix, order_base)
        legacy = os.path.join(output_folder, f"{prefix}_{order_base}.xlsx")
        return path if os.path.isfile(path) or not os.path.isfile(legacy) else legacy

    paths = {
        "calculation": os.path.join(output_folder, f"Розрахунок_розсилки_{order_base}.xlsx"),
        "unmatched": control("Контроль_пропущених_пунктів"),
        "routing": control("Контроль_маршрутизації"),
    }
    errors = []

    def read(key: str) -> list[tuple]:
        try:
            return _saved_xlsx_rows(paths[key])
        except Exception as error:
            errors.append(f"{os.path.basename(paths[key])}: {error}")
            return []

    calculation_source = read("calculation")
    unmatched_source = read("unmatched")
    routing_source = read("routing")
    calculation_rows = [tuple(row[:3]) for row in calculation_source]
    routing_rows = [
        (
            row[0] if len(row) > 0 else "",
            row[1] if len(row) > 1 else "",
            row[2] if len(row) > 2 else "",
            row[5] if len(row) > 5 else (row[-1] if row else ""),
        )
        for row in routing_source
    ]
    # Старі й нові контрольні файли містять правило «зміна до управління».
    # Саме з нього відновлюємо зелені інформаційні рядки, не змінюючи
    # еталонний XLSX розсилки й не запускаючи маршрутизатор повторно.
    management_rows = [
        (MANAGEMENT_RESULT_LABEL, str(row[0]), 1)
        for row in routing_source
        if len(row) > 2 and "управлін" in str(row[2]).casefold()
    ]
    calculation_rows.extend(management_rows)
    return {
        "calculation": calculation_rows,
        "unmatched": [tuple(row[:3]) for row in unmatched_source],
        "routing": routing_rows,
        "found": {key: os.path.isfile(path) for key, path in paths.items()},
        "paths": paths,
        "errors": errors,
    }
