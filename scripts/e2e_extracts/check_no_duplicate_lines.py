"""Жоден рядок наказу не потрапляє у витяг двічі.

Проганяє СПРАВЖНІЙ `App.run_extracts` на наказі з пунктом «ВИКЛАСТИ В ТАКІЙ
РЕДАКЦІЇ» (`make_fixtures_redaction.py`) і перевіряє готовий документ на
повтори абзаців.

**Маршрутизація тут навмисно підмінюється.** На вигаданому наказі вона межі
розставляє чисто, а в бойовому — ні: пункт, процитований усередині лапок,
дістає «шапку», яка починається на рядках самого вступного пункту. Саме через
це у витягу двічі друкувався пункт «11.». Перевіряти треба поведінку ВСТАВКИ
за таких меж, тож перекриття задається явно.

    python scripts/e2e_extracts/make_fixtures_redaction.py <тека>
    python scripts/e2e_extracts/check_no_duplicate_lines.py <тека> .

Код виходу 1 — знайдено повтор.
"""
import glob
import os
import subprocess
import sys

sys.stdout.reconfigure(encoding="utf-8")

E2E = sys.argv[1]
PROJECT = os.path.abspath(sys.argv[2] if len(sys.argv) > 2 else ".")
os.chdir(PROJECT)
sys.path.insert(0, os.path.join(PROJECT, "src"))

import docx  # noqa: E402
from importlib import util  # noqa: E402

spec = util.spec_from_file_location("gen_dup_check", os.path.join(PROJECT, "generate_extracts.py"))
gen = util.module_from_spec(spec)
spec.loader.exec_module(gen)


class _Silent:
    @staticmethod
    def showinfo(*a, **k): pass
    @staticmethod
    def showwarning(*a, **k): pass
    @staticmethod
    def showerror(*a, **k): print("[error]", (a[1] if len(a) > 1 else "")[:600])
    @staticmethod
    def askyesno(*a, **k): return False


gen.messagebox = _Silent
gen.os.startfile = lambda *a, **k: None

_original_map = gen.map_military_units


def _map_with_overlap(**kwargs):
    """Пункту з цитати дописує шапку, яка заходить у вступний пункт."""
    result = _original_map(**kwargs)
    for data in result.get("unit_paragraphs", {}).values():
        items = sorted(data.get("items", []), key=lambda item: item.get("source_start_line", 10 ** 9))
        if len(items) < 2:
            continue
        intro, quoted = items[0], items[1]
        quoted["heading_ranges"] = [
            (intro.get("source_start_line"), quoted.get("source_start_line") - 1)
        ]
        print(
            f"[підміна] {quoted.get('label')}: шапка {quoted['heading_ranges'][0]}, "
            f"вступний пункт {intro.get('source_start_line')}–{intro.get('source_end_line')}"
        )
    return result


gen.map_military_units = _map_with_overlap


class Var:
    def __init__(self, value=""): self._v = value
    def get(self): return self._v
    def set(self, value): self._v = value


order = glob.glob(os.path.join(E2E, "Наказ*.docx"))[0]
out_folder = os.path.join(E2E, "Extracts_Output")

app = gen.App.__new__(gen.App)
app.log = lambda message="": print(message)
app.show_analysis_results = lambda *a, **k: None
app.show_layout_warnings = lambda *a, **k: None
app.excel_path = Var(os.path.join(E2E, "mapping.xlsx"))
app.doc_path = Var(order)
app.template_path = Var(os.path.join(E2E, "template.docx"))
app.out_folder = Var(out_folder)
app.executor = Var("Тест Тестенко 00-000")
app.certifier_position = Var("Начальник відділу /штабу військової частини А0001")
app.certifier_rank = Var("підполковник")
app.certifier_name = Var("Іван ЗАСВІДЧУВАЧ")
app.order_signer_position = Var("")
app.order_signer_rank = Var("")
app.order_signer_name = Var("")
app.group_corps_var = Var(True)
app.duplex_2up_layout = Var(True)


def word_pids():
    out = subprocess.run(
        ["powershell", "-NoProfile", "-Command",
         "Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object { $_.Id }"],
        capture_output=True, text=True)
    return {line.strip() for line in out.stdout.splitlines() if line.strip()}


before = word_pids()
try:
    app.run_extracts()
finally:
    # Закриваємо ЛИШЕ ті екземпляри Word, які зʼявилися під час прогону.
    for pid in word_pids() - before:
        subprocess.run(["powershell", "-NoProfile", "-Command", f"Stop-Process -Id {pid} -Force"],
                       capture_output=True)

result_paths = sorted(glob.glob(os.path.join(out_folder, "Витяги*.docx")))
if not result_paths:
    print("НЕ ЗНАЙДЕНО згенерований документ у", out_folder)
    raise SystemExit(1)

document = docx.Document(result_paths[0])
seen = {}
duplicates = []
for index, paragraph in enumerate(document.paragraphs, start=1):
    text = paragraph.text.strip()
    # Порожні абзаци та короткі службові рядки повторюватись МАЮТЬ.
    if len(text) < 40:
        continue
    if text in seen:
        duplicates.append((seen[text], index, text))
    else:
        seen[text] = index

print("-" * 72)
print("Абзаців у документі:", len(document.paragraphs))
if duplicates:
    print("ПРОБЛЕМИ: той самий текст перенесено двічі")
    for first, second, text in duplicates:
        print(f"  ✗ абзаци {first} і {second}: {text[:90]}…")
    raise SystemExit(1)
print("ПЕРЕВІРКА ПРОЙДЕНА: жоден рядок наказу не потрапив у витяг двічі.")
