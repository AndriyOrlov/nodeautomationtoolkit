"""Варіант шаблону витягу, у якому підписант і засвідчувач зверстані ТАБЛИЦЕЮ.

Замінює `template.docx` у теці фікстур (сам наказ і словник лишаються ті самі,
їх робить `make_fixtures.py`). Усе вигадане, усе на C:.

Навіщо окремий шаблон: у бойовому зразку блок «підписант / Згідно з оригіналом
/ засвідчувач» — це таблиця 5×2, а `{{виконавець}}` стоїть ОДРАЗУ ПІСЛЯ неї,
останнім абзацом документа. Така верстка ловить дві помилки, яких не видно на
шаблоні з самих абзаців:

* опускання виконавця донизу сторінки впиралось у порожню комірку таблиці —
  її знак абзацу Word не видаляє й помилки не кидає, тож цикл був вічним і
  програма зависала на ПЕРШОМУ ж витягу;
* видалення рядків підписанта за наперед зібраним списком абзаців зносило
  ще й сусідній рядок «Згідно з оригіналом».

Запуск:

    python scripts/e2e_extracts/make_fixtures.py <тека>
    python scripts/e2e_extracts/make_template_table_signer.py <тека>
    python scripts/e2e_extracts/run_extracts_e2e.py <тека> .
"""
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

OUT = sys.argv[1]
C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY
R = WD_ALIGN_PARAGRAPH.RIGHT


def para(doc, text="", align=None, size=14, bold=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    return p


def cell_text(cell, text, size=14):
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


tpl = docx.Document()
s = tpl.sections[0]
s.left_margin, s.right_margin = Cm(2.0), Cm(1.0)
s.top_margin, s.bottom_margin = Cm(1.0), Cm(1.0)
normal = tpl.styles["Normal"]
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.alignment = J

para(tpl, "{{кому}}", R)
para(tpl, "{{куди}}", R)
para(tpl)
para(tpl, "ВИТЯГ З НАКАЗУ", C, bold=True)
para(tpl, "КОМАНДИРА ВІЙСЬКОВОЇ ЧАСТИНИ А0001", C, bold=True)
para(tpl)
para(tpl, "(по особовому складу)", C)
para(tpl)

# Дата, місто й номер — теж у таблиці, як у бойовому зразку.
head = tpl.add_table(rows=1, cols=3)
cell_text(head.cell(0, 0), "{{дата_наказу}}")
cell_text(head.cell(0, 1), "м. Тестове")
cell_text(head.cell(0, 2), "{{номер_наказу}}")

para(tpl)
para(tpl, "{{зміст}}")
para(tpl)
para(tpl)

# Підписант, засвідчення й засвідчувач — одна таблиця 5×2.
sign = tpl.add_table(rows=5, cols=2)
for row_index, (left, right) in enumerate([
    ("{{підписант_посада}}", ""),
    ("{{підписант_звання}}", "{{підписант_піб}}"),
    ("Згідно з оригіналом", ""),
    ("{{засвідчувач_посада}}", ""),
    ("{{засвідчувач_звання}}", "{{Засвідчувач}}"),
]):
    cell_text(sign.cell(row_index, 0), left)
    cell_text(sign.cell(row_index, 1), right)

# Виконавець — ОСТАННІЙ абзац документа, одразу за таблицею.
executor = tpl.add_paragraph()
executor_run = executor.add_run("{{виконавець}}")
executor_run.font.name = "Times New Roman"
executor_run.font.size = Pt(8)

path = os.path.join(OUT, "template.docx")
tpl.save(path)
print("шаблон із табличним підписантом:", path)
