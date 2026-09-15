"""Перевірка згенерованого повідомлення: шрифт і правила розд. 9.

Працює ВИКЛЮЧНО на вигаданих фікстурах (`make_message_fixtures.py`), тому
вміст можна виводити: персональних даних у ньому немає.
"""
import glob
import os
import re
import sys

import docx

sys.stdout.reconfigure(encoding="utf-8")

OUT = sys.argv[1]
path = sorted(glob.glob(os.path.join(OUT, "Повідомлення_шифрований_зміст*.docx")))
if not path:
    print("НЕ ЗНАЙДЕНО файл зі змістом у", OUT)
    raise SystemExit(1)
document = docx.Document(path[0])
normal_font = document.styles["Normal"].font.name


def effective_font(run, paragraph):
    for candidate in (
        run.font.name,
        run.style.font.name if run.style is not None else None,
        paragraph.style.font.name if paragraph.style is not None else None,
        normal_font,
    ):
        if candidate:
            return candidate
    return "(невідомо)"


def is_order_content(text: str) -> bool:
    """Абзац, перенесений з наказу (а не власний текст шаблону).

    Шрифт перевіряємо лише для нього: власні рядки зразка лишаються такими,
    як їх набрав користувач (правило 9.4), і в фікстурі вони навмисно Calibri.
    """
    clean = (text or "").strip()
    if not clean:
        return False
    return bool(
        clean.startswith("§")
        or re.match(r"^\d{1,3}(?:\.\d{1,3})*[.)]\s", clean)  # 1. і 1.1.
        or clean.startswith("Відповідно до")
        or re.search(r"р\.\s?н\.|у ЗС|Підлягає направленню", clean)
        or re.match(r"^\d{10}\.?$", clean)
    )


fonts = {}
template_fonts = {}
lines = []
for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    target = fonts if is_order_content(text) else template_fonts
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        name = effective_font(run, paragraph)
        target[name] = target.get(name, 0) + len(run.text)
    if text:
        lines.append(text)

body = "\n".join(lines)
problems = []

# Пункт 1.2 фікстури: половина «КУДИ» має лишитися цілою. Саме її
# з'їдав один збіг словника, що перетинав тире-роздільник (4.2.9).
if "НАЧАЛЬНИКОМ ГРУПИ ПСИХОЛОГІЧНОЇ ПІДТРИМКИ ПЕРСОНАЛУ" not in body:
    problems.append("ВТРАТА ТЕКСТУ: зникла половина «КУДИ» пункту 1.2")

if set(fonts) - {"Times New Roman"}:
    problems.append(f"НЕ Times New Roman: {fonts}")

for pattern, why in (
    (r"частин\w+\s+військов", "дубль «військової частини військової частини»"),
    (r"\bокрем\w*\b", "відкрита назва частини лишилась"),
    (r"\bбригад\w*\b", "відкрита назва частини лишилась"),
    (r"\bкорпус\w*\b", "відкрита назва корпусу лишилась"),
    (r"[«“][^»”]{0,40}Яр[»”]", "почесне найменування в лапках лишилось"),
    (r"\bАК\b", "скорочення корпусу зі стовпця D потрапило в текст"),
    (r"\b[РМО]?ТЦК\b", "назву ТЦК стиснуто до короткої форми (має лишатись повною)"),
    (r"(військов\w+\s+частин\w+\s+[АA]\d+)\s+\1\b", "той самий шифр двічі підряд"),
):
    for match in re.finditer(pattern, body, re.IGNORECASE):
        problems.append(f"{why}: …{body[max(0, match.start() - 45):match.end() + 25]}…")

# Супровід: у колонтитулі теж не має лишитись незаповнених тегів. Колонтитул —
# окрема «історія» Word, тому заміна тегів у тілі документа його не зачіпає.
cover_paths = sorted(glob.glob(os.path.join(OUT, "Повідомлення_супровід*.docx")))
if cover_paths:
    cover = docx.Document(cover_paths[0])
    for section_index, section in enumerate(cover.sections, 1):
        for part_name, part in (("нижній", section.footer), ("верхній", section.header)):
            for paragraph in part.paragraphs:
                if "{{" in paragraph.text:
                    problems.append(
                        f"{part_name} колонтитул супроводу (розділ {section_index}): "
                        f"тег лишився незаповненим — {paragraph.text.strip()!r}"
                    )

print("Шрифти ЗМІСТУ з наказу (символів):", fonts)
print("Шрифти власних рядків шаблону:", template_fonts)
print("Стиль Normal шаблона:", normal_font)
print("Абзаців із текстом:", len(lines))
print("-" * 72)
for line in lines:
    print("  ", line)
print("-" * 72)
if problems:
    print("ПРОБЛЕМИ:")
    for problem in problems:
        print("  ✗", problem)
    raise SystemExit(1)
print("ПЕРЕВІРКА ПРОЙДЕНА: шрифт Times New Roman, дублів і відкритих назв немає.")
